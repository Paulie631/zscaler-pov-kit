"""The POV's executive summary, in two versions, as Word and PowerPoint files.

  customer  for the customer's executives: what the POV set out to show, what it proved in business terms, the AI
            Guard numbers, what's still open, and a clear recommended next step.
  internal  for Zscaler's deal review: the technical verdict, win themes, risks, gaps to close, competitive notes and
            asks. Marked internal; never sent to the customer.

Claude writes both from the results alone, as the older summary did: criteria and outcomes, the use cases' requests,
the AI Guard test counts, and the SE's own notes, with the customer's names taken out. Claude writes {CUSTOMER} where
the customer is named; the kit puts the real name in on this computer, so the name never leaves it. The files are made
here (python-docx, python-pptx) when the SE downloads them.

The customer version also becomes the summary that opens the evidence report (summary.json), so there's one summary.
"""
from __future__ import annotations

import io
import json
import time
from collections import Counter

from . import ai, evidence, products
from .core import KitError
from .core.redact import Redactor

BLUE, INK, QUIET, OK, BAD, MID = "1F5FD6", "1D2330", "5B6475", "1F7A4D", "B3261E", "946200"

SYSTEM = """You write the executive summary of a Zscaler proof of value (POV), in two versions, from its results.
Write {CUSTOMER} wherever you name the customer (the kit fills in the name). Be plain, specific and honest: only a
criterion whose result is "pass" is proved; "fail" and "incomplete" are not, and say so. Never invent numbers,
products, dates or results; use only the figures given. Short sentences, no jargon an executive wouldn't know.

customer (for {CUSTOMER}'s executives, sent by the Zscaler SE):
  headline: one sentence, what the POV showed overall
  situation: two or three sentences: what {CUSTOMER} wanted to prove and with which Zscaler products
  results: up to 6 bullets, each something shown to work, in business terms, citing criterion IDs in brackets
  ai_guard: two or three sentences on the AI Guard testing with its numbers, or "" if there was none
  value: up to 4 bullets: what this means for {CUSTOMER} (risk reduced, effort saved, users protected), only where
         the results support it
  open: what isn't proved yet and what it would take, citing IDs; [] if nothing
  recommendation: one or two sentences, a clear recommended next step (e.g. move the proved use cases to production),
                  scoped to what was proved
  next_steps: up to 4 short bullets with owners where obvious ({CUSTOMER}, Zscaler)

internal (Zscaler deal review only; candid):
  verdict: one sentence: technical win, partial, or not yet, and why
  win_themes: up to 4 bullets: what resonated or was proved that sells the deal
  risks: up to 4 bullets: technical or deal risks the results show (failures, gaps, untested areas)
  gaps_to_close: up to 4 bullets: what must happen before the decision, citing IDs
  competitive: up to 3 bullets, only from the SE's notes; [] if none given
  recommendation: one or two sentences: what the account team should do next
  asks: up to 4 bullets: help needed (product team, specialists, executive sponsor)

Call write_exec_summary with both."""

_ARR = {"type": "array", "items": {"type": "string"}}
_STR = {"type": "string"}
CUST_KEYS = {"headline": _STR, "situation": _STR, "results": _ARR, "ai_guard": _STR, "value": _ARR, "open": _ARR,
             "recommendation": _STR, "next_steps": _ARR}
INT_KEYS = {"verdict": _STR, "win_themes": _ARR, "risks": _ARR, "gaps_to_close": _ARR, "competitive": _ARR,
            "recommendation": _STR, "asks": _ARR}
TOOL = {"name": "write_exec_summary", "description": "The POV's executive summary, customer and internal versions.",
        "input_schema": {"type": "object", "required": ["customer", "internal"], "properties": {
            "customer": {"type": "object", "required": list(CUST_KEYS), "properties": CUST_KEYS},
            "internal": {"type": "object", "required": list(INT_KEYS), "properties": INT_KEYS}}}}


# ---------------------------------------------------------------- what Claude sees
def _redactor(pov):
    snap = (pov.read_json("discovery.json", {}) or {}).get("products") or {}
    names = {pov.cfg.get("customer") or "", pov.cfg.get("pilot_group") or "", pov.cfg.get("vanity_domain") or ""}
    for p in products.selected(pov.products):
        if p.name in snap:
            names |= p.customer_names(snap[p.name])
    return Redactor(names, pov.prefix)


def aiguard_stats(pov):
    """The AI Guard test runner's numbers: scans, flagged, blocked, the detections and models. None without runs."""
    from .aiguard_runner import PANELS, runs
    rs = [r for r in runs(pov) if r.get("status") == "ok"]
    if not rs:
        return None
    det = Counter(d for r in rs for d in r.get("detectors") or [])
    attacks = [r for r in rs if r.get("expect")]
    caught = [r for r in attacks if r.get("pass") is True]
    return {"scans": len(rs), "flagged": sum(1 for r in rs if r.get("detectors")),
            "blocked": sum(1 for r in rs if r.get("action") == "BLOCK"),
            "harmful_tests": len(attacks), "harmful_caught": len(caught),
            "detections": [{"type": PANELS.get(d, d), "count": n} for d, n in det.most_common(8)],
            "models": sorted({r["provider"] for r in rs if r.get("provider")}),
            "modes": sorted({r.get("mode") or "detect" for r in rs})}


def facts(pov):
    red = _redactor(pov)
    rows = evidence.evaluate(pov)
    try:
        from . import usecases
        ucs = [red.text(u.get("request") or (u.get("test") or {}).get("title") or "") for u in usecases.accepted(pov)]
    except Exception:   # noqa: BLE001
        ucs = []
    count = Counter(r["status"] for r in rows)
    return {"products": [products.get(p).title for p in pov.products],
            "started": pov.cfg.get("created", "")[:10], "ends": pov.cfg.get("ends"),
            "use_cases": [u for u in ucs if u][:12],
            "criteria_totals": {"total": len(rows), "pass": count.get("pass", 0), "fail": count.get("fail", 0),
                                "incomplete": count.get("incomplete", 0)},
            "criteria": [{"id": r["criterion"]["id"], "title": red.text(r["criterion"]["title"]),
                          "product": r["criterion"]["product"], "result": r["status"],
                          "proofs": {k: red.text(str(v.get("detail") or v.get("status")))[:240]
                                     for k, v in r["proofs"].items()}} for r in rows],
            "ai_guard": aiguard_stats(pov)}, red


def _clean(d, keys):
    out = {}
    for k, t in keys.items():
        v = (d or {}).get(k)
        out[k] = ([str(x)[:400] for x in (v or []) if str(x).strip()][:6] if t is _ARR else str(v or "")[:700])
    return out


def write(pov, emphasis="", internal_notes="", prepared_by=""):
    """Ask Claude for both versions; saved as execsum.json (and the report's summary.json)."""
    f, red = facts(pov)
    if not f["criteria"]:
        raise KitError("No success criteria yet: choose them first, so there's something to summarize.")
    msg = "POV results (JSON):\n" + json.dumps(f)
    if emphasis.strip():
        msg += f"\n\nThe SE asks you to emphasize: {red.text(emphasis.strip())[:600]}"
    if internal_notes.strip():
        msg += f"\n\nThe SE's internal notes (deal context, competitors; internal version only): " \
               f"{red.text(internal_notes.strip())[:1500]}"
    used, text = ai.ask(pov, SYSTEM, msg, max_tokens=4000, tool=TOOL, needs_consent=False)
    if used is None:
        from .usecases import parse
        used = parse(text)
    now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    out = {"customer": _clean(used.get("customer"), CUST_KEYS), "internal": _clean(used.get("internal"), INT_KEYS),
           "facts": {k: f[k] for k in ("products", "started", "ends", "criteria_totals", "ai_guard")},
           "criteria": [{k: c[k] for k in ("id", "title", "result")} for c in f["criteria"]],
           "results": {c["id"]: c["result"] for c in f["criteria"]}, "at": now,
           "prepared_by": prepared_by.strip()[:120], "emphasis": emphasis.strip()[:600]}
    pov.write_json("execsum.json", out)
    cu = out["customer"]
    pov.write_json("summary.json", {"headline": cu["headline"], "proved": cu["results"], "open": cu["open"],
                                    "value": cu["value"], "next_steps": cu["next_steps"], "at": now,
                                    "results": out["results"]})
    pov.log("report", f"executive summary (customer and internal) written by Claude {ai.ROUTES[ai.route_name()]}")
    return out


def current(pov):
    """The saved summary, and whether the results changed since."""
    s = pov.read_json("execsum.json")
    if not s:
        return None, False
    now = {r["criterion"]["id"]: r["status"] for r in evidence.evaluate(pov)}
    return s, now != s.get("results")


def fill(pov, s):
    """{CUSTOMER} -> the customer's name, everywhere (only on this computer)."""
    name = pov.cfg.get("customer") or "the customer"

    def go(x):
        if isinstance(x, str):
            return x.replace("{CUSTOMER}", name)
        if isinstance(x, list):
            return [go(y) for y in x]
        if isinstance(x, dict):
            return {k: go(v) for k, v in x.items()}
        return x
    return go(s)


def filename(pov, version, fmt):
    import re
    who = re.sub(r"[^A-Za-z0-9]+", "-", pov.cfg.get("customer") or pov.id).strip("-") or "POV"
    tag = "Executive-Summary" if version == "customer" else "Internal-Readout"
    return f"{who}-POV-{tag}.{fmt}"


def render(pov, version, fmt):
    if version not in ("customer", "internal") or fmt not in ("docx", "pptx"):
        raise KitError("Choose the customer or internal version, as docx or pptx.")
    s, _ = current(pov)
    if not s:
        raise KitError("Draft the executive summary first.")
    s = fill(pov, s)
    try:
        return (_docx if fmt == "docx" else _pptx)(pov, s, version)
    except ImportError as e:
        raise KitError("This kit is missing python-docx or python-pptx: restart the console (it fetches them), or "
                       "update the kit.") from e


# ---------------------------------------------------------------- shared wording
RESULT = {"pass": "Proved", "fail": "Not proved", "incomplete": "Still open"}


def _meta(pov, s):
    f = s["facts"]
    bits = [", ".join(f["products"]), f"{f['started']} to {f.get('ends') or 'ongoing'}"]
    if s.get("prepared_by"):
        bits.append(f"Prepared by {s['prepared_by']}")
    return " · ".join(b for b in bits if b)


def _kpis(s):
    t, ag = s["facts"]["criteria_totals"], s["facts"].get("ai_guard")
    k = [(f"{t['pass']} of {t['total']}", "success criteria proved")]
    if ag:
        k.append((str(ag["scans"]), "AI prompts and replies tested"))
        if ag["harmful_tests"]:
            k.append((f"{ag['harmful_caught']} of {ag['harmful_tests']}", "risky prompts caught"))
        k.append((str(ag["blocked"]), "blocked by AI Guard"))
    return k


def _sections(s, version):
    if version == "customer":
        c = s["customer"]
        return [("Summary", [c["headline"]], False), ("What we set out to prove", [c["situation"]], False),
                ("What the POV proved", c["results"], True), ("AI Guard testing", [c["ai_guard"]] if c["ai_guard"] else [], False),
                ("What it means for the business", c["value"], True), ("Still open", c["open"], True),
                ("Our recommendation", [c["recommendation"]], False), ("Next steps", c["next_steps"], True)]
    i = s["internal"]
    return [("Verdict", [i["verdict"]], False), ("Recommendation", [i["recommendation"]], False),
            ("Win themes", i["win_themes"], True), ("Risks", i["risks"], True),
            ("Gaps to close before the decision", i["gaps_to_close"], True), ("Competitive", i["competitive"], True),
            ("Asks", i["asks"], True)]


# ---------------------------------------------------------------- Word
def _docx(pov, s, version):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    def rgb(h):
        return RGBColor.from_string(h)

    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(60)
    name = pov.cfg.get("customer") or "Customer"
    if version == "internal":
        b = d.add_paragraph()
        r = b.add_run("ZSCALER INTERNAL · NOT FOR THE CUSTOMER")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), rgb(BAD)
    top = d.add_paragraph()
    r = top.add_run("Zscaler · Proof of Value")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(10), rgb(BLUE)
    h = d.add_paragraph()
    r = h.add_run(f"{name}: " + ("Executive Summary" if version == "customer" else "Internal POV Readout"))
    r.bold, r.font.size, r.font.color.rgb = True, Pt(22), rgb(INK)
    m = d.add_paragraph()
    r = m.add_run(_meta(pov, s))
    r.font.size, r.font.color.rgb = Pt(9.5), rgb(QUIET)

    kp = _kpis(s)
    t = d.add_table(rows=2, cols=len(kp))
    for n, (big, small) in enumerate(kp):
        c1, c2 = t.cell(0, n), t.cell(1, n)
        r = c1.paragraphs[0].add_run(big)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(20), rgb(BLUE)
        r = c2.paragraphs[0].add_run(small)
        r.font.size, r.font.color.rgb = Pt(9), rgb(QUIET)
        for c in (c1, c2):
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def heading(text):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(13), rgb(BLUE)

    for title, items, bullets in _sections(s, version):
        if not [x for x in items if x]:
            continue
        heading(title)
        for x in items:
            p = d.add_paragraph(x, style="List Bullet" if bullets else None)
            if title in ("Summary", "Verdict"):
                for r in p.runs:
                    r.bold = True

    heading("Success criteria")
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    for c, txt in zip(t.rows[0].cells, ("ID", "Criterion", "Result"), strict=True):
        c.paragraphs[0].add_run(txt).bold = True
    for c in s["criteria"]:
        row = t.add_row().cells
        row[0].text, row[1].text = c["id"], c["title"]
        r = row[2].paragraphs[0].add_run(RESULT.get(c["result"], c["result"]))
        r.bold, r.font.color.rgb = True, rgb({"pass": OK, "fail": BAD}.get(c["result"], MID))
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    from docx.shared import Inches
    t.autofit = False
    for row in t.rows:
        for c, w in zip(row.cells, (Inches(1.1), Inches(4.3), Inches(1.1)), strict=True):
            c.width = w
    foot = d.add_paragraph()
    foot.paragraph_format.space_before = Pt(18)
    r = foot.add_run(("Drafted with Claude from the POV's recorded results; reviewed by the Zscaler SE. "
                      if version == "customer" else "Drafted with Claude from the POV's results and the SE's notes. ")
                     + f"Results as of {s['at'][:10]}.")
    r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), rgb(QUIET)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PowerPoint
def _pptx(pov, s, version):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    W, H = Emu(12192000), Emu(6858000)   # noqa: N806 - 16:9
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    name = pov.cfg.get("customer") or "Customer"
    internal = version == "internal"
    M = Emu(560000)   # noqa: N806

    def rgb(h):
        return RGBColor.from_string(h)

    def box(sl, x, y, w, h, text="", size=16, bold=False, color=INK, fill=None):
        sh = sl.shapes.add_textbox(x, y, w, h) if not fill else sl.shapes.add_shape(1, x, y, w, h)
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
            sh.line.fill.background()
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text, r.font.size, r.font.bold, r.font.color.rgb = text, Pt(size), bold, rgb(color)
        return tf

    def frame(title):
        sl = prs.slides.add_slide(blank)
        box(sl, 0, 0, W, Emu(90000), fill=BLUE)
        box(sl, M, Emu(260000), W - 2 * M, Emu(700000), title, 28, True)
        tag = "Zscaler POV · " + name + (" · INTERNAL" if internal else "")
        box(sl, M, H - Emu(480000), W - 2 * M, Emu(300000), tag, 10, False, BAD if internal else QUIET)
        return sl

    def bullets(sl, items, x, y, w, h, size=17):
        tf = box(sl, x, y, w, h, "", size)
        first = True
        for it in [i for i in items if i]:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(8)
            r = p.add_run()
            r.text, r.font.size, r.font.color.rgb = "•  " + it, Pt(size), rgb(INK)

    # title slide
    sl = prs.slides.add_slide(blank)
    box(sl, 0, 0, W, H, fill=BLUE)
    if internal:
        box(sl, M, Emu(500000), W - 2 * M, Emu(400000), "ZSCALER INTERNAL · NOT FOR THE CUSTOMER", 12, True, "FFD7D4")
    box(sl, M, Emu(2100000), W - 2 * M, Emu(500000), "Zscaler · Proof of Value", 18, True, "DCE6FB")
    box(sl, M, Emu(2600000), W - 2 * M, Emu(1100000),
        f"{name}: " + ("Executive Summary" if not internal else "Internal POV Readout"), 40, True, "FFFFFF")
    box(sl, M, Emu(3900000), W - 2 * M, Emu(500000), _meta(pov, s), 14, False, "DCE6FB")

    kp = _kpis(s)
    if not internal:
        c = s["customer"]
        sl = frame("Results at a glance")
        box(sl, M, Emu(1050000), W - 2 * M, Emu(800000), c["headline"], 20, True)
        n = len(kp)
        gap, kw = Emu(200000), int((W - 2 * M - Emu(200000) * (n - 1)) / n)
        for i, (big, small) in enumerate(kp):
            x = M + i * (kw + gap)
            box(sl, x, Emu(2200000), kw, Emu(1500000), fill="EEF3FD")
            box(sl, x + Emu(150000), Emu(2350000), kw - Emu(300000), Emu(700000), big, 36, True, BLUE)
            box(sl, x + Emu(150000), Emu(3100000), kw - Emu(300000), Emu(500000), small, 13, False, QUIET)
        box(sl, M, Emu(4000000), W - 2 * M, Emu(1200000), c["situation"], 15, False, QUIET)
        sl = frame("What the POV proved")
        bullets(sl, c["results"], M, Emu(1100000), W - 2 * M, Emu(4300000))
        if c["ai_guard"] or c["value"]:
            sl = frame("AI Guard and what it means" if c["ai_guard"] else "What it means for the business")
            y = Emu(1100000)
            if c["ai_guard"]:
                box(sl, M, y, W - 2 * M, Emu(1100000), c["ai_guard"], 16, False)
                y = Emu(2350000)
            bullets(sl, c["value"], M, y, W - 2 * M, H - y - Emu(700000))
        sl = frame("Recommendation and next steps")
        box(sl, M, Emu(1050000), W - 2 * M, Emu(900000), c["recommendation"], 20, True, BLUE)
        half = int((W - 2 * M - Emu(300000)) / 2)
        box(sl, M, Emu(2150000), half, Emu(400000), "Next steps", 14, True, QUIET)
        bullets(sl, c["next_steps"], M, Emu(2550000), half, Emu(2900000), 15)
        if c["open"]:
            box(sl, M + half + Emu(300000), Emu(2150000), half, Emu(400000), "Still open", 14, True, QUIET)
            bullets(sl, c["open"], M + half + Emu(300000), Emu(2550000), half, Emu(2900000), 15)
    else:
        i = s["internal"]
        sl = frame("Verdict")
        box(sl, M, Emu(1050000), W - 2 * M, Emu(900000), i["verdict"], 22, True)
        box(sl, M, Emu(2100000), W - 2 * M, Emu(900000), i["recommendation"], 17, False, BLUE)
        line = "   ·   ".join(f"{b} {sm}" for b, sm in kp)
        box(sl, M, Emu(3300000), W - 2 * M, Emu(600000), line, 14, False, QUIET)
        half = int((W - 2 * M - Emu(300000)) / 2)
        for title, (la, a), (lb, b) in (("Win themes and risks", ("Win themes", i["win_themes"]), ("Risks", i["risks"])),
                                        ("Gaps, competition and asks", ("Gaps to close", i["gaps_to_close"]),
                                         ("Asks", i["asks"] + [f"Competitive: {x}" for x in i["competitive"]]))):
            sl = frame(title)
            for k, (lab, items) in enumerate(((la, a), (lb, b))):
                x = M + k * (half + Emu(300000))
                box(sl, x, Emu(1100000), half, Emu(400000), lab, 15, True, QUIET)
                bullets(sl, items, x, Emu(1500000), half, Emu(3900000), 15)

    # the criteria, on every version
    sl = frame("Success criteria")
    rows = s["criteria"][:14]
    tb = sl.shapes.add_table(len(rows) + 1, 3, M, Emu(1050000), W - 2 * M, Emu(300000) * (len(rows) + 1)).table
    tb.columns[0].width, tb.columns[2].width = Emu(1500000), Emu(1700000)
    tb.columns[1].width = W - 2 * M - Emu(3200000)
    for j, txt in enumerate(("ID", "Criterion", "Result")):
        tb.cell(0, j).text = txt
    for r_, c in enumerate(rows, 1):
        tb.cell(r_, 0).text, tb.cell(r_, 1).text = c["id"], c["title"]
        tb.cell(r_, 2).text = RESULT.get(c["result"], c["result"])
        tb.cell(r_, 2).text_frame.paragraphs[0].runs[0].font.color.rgb = rgb({"pass": OK, "fail": BAD}.get(c["result"], MID))
    for row in tb.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(11)
    if len(s["criteria"]) > 14:
        box(sl, M, H - Emu(800000), W - 2 * M, Emu(300000), f"…and {len(s['criteria']) - 14} more in the evidence report.",
            11, False, QUIET)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
